from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "VIC_Cipher_Team6_Summary_Report.docx"


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(23, 32, 51)
    return heading


def add_body_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(35, 42, 55)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(10.2)
        run.font.color.rgb = RGBColor(35, 42, 55)


def build_report():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("VIC Cipher / Spy App")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(17, 24, 39)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("Week 12 Cryptography Summary Report | Team 6 Individual Work")
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(90, 100, 120)

    info = doc.add_table(rows=3, cols=2)
    info.autofit = True
    rows = [
        ("Student Name / ID", "______________________________"),
        ("Topic", "VIC Cipher / Spy App"),
        ("Deliverables", "Summary report, PPT presentation, Python source-code demo, browser demo"),
    ]
    for row, (left, right) in zip(info.rows, rows):
        set_cell_text(row.cells[0], left, bold=True, color=(37, 99, 235))
        set_cell_text(row.cells[1], right)

    doc.add_paragraph()

    add_heading(doc, "1. Executive Summary", 1)
    add_body_paragraph(
        doc,
        "This report summarizes the VIC Cipher / spy app topic for Week 12 Cryptography. "
        "The VIC cipher is a Cold War-era hand cipher associated with covert communication. "
        "It is important because it demonstrates how strong cryptographic ideas existed before "
        "modern computers. The project explains the cipher using W5H1, demonstrates a "
        "VIC-inspired Python implementation, and analyzes possible cracking methods and ethical limits.",
    )

    add_heading(doc, "2. W5H1 Analysis", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.rows[0].cells[0], "Question", bold=True, color=(37, 99, 235))
    set_cell_text(table.rows[0].cells[1], "Summary", bold=True, color=(37, 99, 235))
    w5h1 = [
        ("Who", "Cold War intelligence users; the VIC cipher is connected to Soviet spy Reino Hayhanen."),
        ("What", "A hand cipher that protects plaintext by converting it into encrypted number groups."),
        ("When", "Cold War period; publicly known after espionage investigations."),
        ("Where", "Covert communication settings where messages could be intercepted."),
        ("Why", "To protect secret messages without requiring computers or electronic devices."),
        ("How", "Through secret key material, numeric substitution, transposition, and careful message handling."),
    ]
    for left, right in w5h1:
        cells = table.add_row().cells
        set_cell_text(cells[0], left, bold=True)
        set_cell_text(cells[1], right)

    add_heading(doc, "3. Cryptography Concepts", 1)
    add_bullets(
        doc,
        [
            "Plaintext: the original readable message.",
            "Ciphertext: the encrypted message sent to another person.",
            "Key: secret information required to encrypt and decrypt correctly.",
            "Substitution: replacing letters with other symbols, letters, or numbers.",
            "Transposition: rearranging the order of symbols to hide patterns.",
        ],
    )

    add_heading(doc, "4. Demo Design", 1)
    add_body_paragraph(
        doc,
        "The demo is VIC-inspired, not a full historical VIC cipher implementation. The full historical "
        "cipher is complex and designed for manual use. The project demo focuses on the important "
        "classroom ideas: a secret key controls the encryption process, plaintext becomes ciphertext, "
        "and the correct key is needed to recover the original message.",
    )
    add_bullets(
        doc,
        [
            "Source code file: vic_cipher_professional_demo.py",
            "Visual browser demo: vic_spy_console.html",
            "Input example: plaintext MEET AT NOON and key SECRET 1970",
            "Main techniques: keyed alphabet, straddling checkerboard substitution, and columnar transposition",
        ],
    )

    add_heading(doc, "5. Demo Result", 1)
    result_table = doc.add_table(rows=1, cols=2)
    result_table.style = "Table Grid"
    set_cell_text(result_table.rows[0].cells[0], "Step", bold=True, color=(37, 99, 235))
    set_cell_text(result_table.rows[0].cells[1], "Result", bold=True, color=(37, 99, 235))
    demo_rows = [
        ("Plaintext", "MEET AT NOON"),
        ("Normalized plaintext", "MEETATNOON"),
        ("Secret key", "SECRET 1970"),
        ("Ciphertext", "00151 88722 59012 02597 20"),
        ("Correct-key decryption", "MEETATNOON"),
        ("Wrong-key test", "Does not recover the original plaintext"),
    ]
    for left, right in demo_rows:
        cells = result_table.add_row().cells
        set_cell_text(cells[0], left, bold=True)
        set_cell_text(cells[1], right)

    add_heading(doc, "6. Cracking Possibilities", 1)
    add_body_paragraph(
        doc,
        "Classical ciphers can be attacked when the key is weak, repeated, or used incorrectly. "
        "The VIC cipher was strong for a hand cipher, but any manual cryptographic system can become "
        "weaker if the user makes mistakes or repeats patterns.",
    )
    add_bullets(
        doc,
        [
            "Weak or guessable keys can reduce security.",
            "Repeated key material can expose patterns.",
            "Human errors during encryption can reveal structure.",
            "Repeated phrases or predictable message formats can help attackers.",
            "Modern computers can test possibilities faster than human cryptanalysts.",
            "Modern real-world systems should use reviewed algorithms such as AES, RSA, or ECC.",
        ],
    )

    add_heading(doc, "7. Ethics And Limitations", 1)
    add_body_paragraph(
        doc,
        "This project is for education and authorized classroom demonstration only. The demo should "
        "not be used for real-world secure communication. Cryptography is important for privacy, "
        "banking, messaging, and national security, but security claims must be made carefully. "
        "The correct lesson is understanding how keys, substitution, and transposition protect information.",
    )

    add_heading(doc, "8. Conclusion", 1)
    add_body_paragraph(
        doc,
        "The VIC cipher shows that meaningful cryptography existed before modern computers. "
        "The project demonstrates the main idea through a Python source-code demo and visual spy console: "
        "a message can be encrypted into number groups, and the correct key is required to recover it. "
        "The report also explains cracking possibilities and why ethical boundaries matter in information security.",
    )

    add_heading(doc, "9. Short Presentation Script", 1)
    add_body_paragraph(
        doc,
        "My topic is the VIC Cipher / spy app for Week 12 cryptography. I will summarize it using W5H1, "
        "then explain the cryptography terms plaintext, ciphertext, key, substitution, and transposition. "
        "After that I will show my demo. The demo uses a keyed alphabet, checkerboard substitution, and "
        "columnar transposition. With the correct key, the ciphertext decrypts back to the original message. "
        "With the wrong key, it does not recover the plaintext. Finally, I will explain cracking possibilities, "
        "including weak keys, repeated keys, human mistakes, and pattern analysis. This demo is educational, "
        "not real-world secure encryption.",
    )

    doc.add_page_break()
    add_heading(doc, "Submission Checklist", 1)
    add_bullets(
        doc,
        [
            "Summary report: VIC_Cipher_Team6_Summary_Report.docx",
            "Presentation: VIC_Cipher_Team6_FINAL_ANIMATED_STYLE.pptx",
            "Python source code: vic_cipher_professional_demo.py",
            "Browser visual demo: vic_spy_console.html",
            "Demo command guide: VIC_Professional_Demo_Commands.md",
        ],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_report()
